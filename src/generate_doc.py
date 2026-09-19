from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


def generate_report(report, template_path, output_path):

    doc = Document()

    # Page margins
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # Default font
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    # =========================
    # TITLE
    # =========================

    title = doc.add_paragraph()

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run("EQUIPMENT INSPECTION REPORT")
    run.bold = True
    run.font.size = Pt(18)

    subtitle = doc.add_paragraph()

    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = subtitle.add_run(report.title)
    run.bold = True
    run.font.size = Pt(13)

    # =========================
    # DOCUMENT INFORMATION
    # =========================

    doc.add_heading("Document Information", level=1)

    info_table = doc.add_table(rows=2, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.style = "Table Grid"

    info_table.cell(0, 0).text = "Document Type"
    info_table.cell(0, 1).text = report.document_type

    info_table.cell(1, 0).text = "Overall Severity"
    info_table.cell(1, 1).text = report.overall_severity

    # =========================
    # SUMMARY
    # =========================

    doc.add_heading("Summary", level=1)

    paragraph = doc.add_paragraph()
    paragraph.add_run(report.summary)

    # =========================
    # FINDINGS
    # =========================

    doc.add_heading("Findings", level=1)

    table = doc.add_table(
        rows=1,
        cols=4
    )

    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = [
        "Item",
        "Finding",
        "Severity",
        "Recommendation"
    ]

    for i, header in enumerate(headers):

        cell = table.rows[0].cells[i]

        cell.text = header

        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for finding in report.findings:

        row = table.add_row()

        row.cells[0].text = finding.item
        row.cells[1].text = finding.finding
        row.cells[2].text = finding.severity
        row.cells[3].text = finding.recommendation

    # =========================
    # RECOMMENDATIONS
    # =========================

    doc.add_heading("Recommendations", level=1)

    for recommendation in report.recommendations:

        paragraph = doc.add_paragraph(
            style="List Bullet"
        )

        paragraph.add_run(recommendation)

    # =========================
    # FOOTER
    # =========================

    footer = section.footer

    paragraph = footer.paragraphs[0]

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(
        "Generated automatically from document analysis"
    )

    run.font.size = Pt(8)

    # Save
    doc.save(output_path)

    print(f"\nWord report created: {output_path}")